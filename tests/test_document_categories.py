"""Regressions from the supplied equivalences: category boundaries and omissions."""
import io

import pdfplumber
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.anonymize.apply import anonymize_text
from app.anonymize.placeholders import build_detections_from_mentions
from app.detection import pipeline
from app.detection.filters import is_valid_detection
from app.detection.regex_ar import _clean_persona_surface
from app.main import app
from app.models.schemas import Cluster, Detection, Mention, SessionState
from app.services.clusters import confirm_cluster


@pytest.mark.parametrize('text, expected', [
    ('con costas.\n\nFCA Ahorro SA', 'FCA Ahorro SA'),
    ('y FCA Ahorro SA', 'FCA Ahorro SA'),
    ('una empleada de la concesionaria Armada SA', 'Armada SA'),
    ('no media relación contractual entre Armada SA', 'Armada SA'),
    ('por Gerardo Alfredo Donati contra\n\nArmada SA', 'Armada SA'),
    ('en el caso de FCA Ahorro SA', 'FCA Ahorro SA'),
    ('Comercial\n\nSala B\n\nVI. FCA Ahorro SA', 'FCA Ahorro SA'),
    ('el señor Donati adquirió de Armada SA', 'Armada SA'),
    ('lo Comercial\n\nSala B\n\nFCA Ahorro SA', 'FCA Ahorro SA'),
    ('en lo Comercial\n\nSala B\n\nAhorro SA', 'Ahorro SA'),
    ('esta instancia la responsabilidad de\n\nArmada SA', 'Armada SA'),
    ('si cabe responsabilizar a FCA Ahorro SA', 'FCA Ahorro SA'),
    ('que la responsabilidad de Armada SA', 'Armada SA'),
    ('de ahorro administrado por FCA Ahorro SA', 'FCA Ahorro SA'),
    ('el expediente surge que\n\nFCA Ahorro SA', 'FCA Ahorro SA'),
    ('confirmar la responsabilidad de FCA Ahorro SA', 'FCA Ahorro SA'),
    ('rechaza el agravio de FCA Ahorro SA', 'FCA Ahorro SA'),
    ('Sala B\n\nTurismo Noche y Dia SRL', 'Turismo Noche y Dia SRL'),
    ('las de Alzada a FCA Ahorro SA', 'FCA Ahorro SA'),
    ('CONFIRMAR LA RESPONSABILIDAD DE FCA AHORRO SA', 'FCA AHORRO SA'),
    ('Autofrance\n\nSA', 'Autofrance\n\nSA'),
    ('Fiat\n\nAuto SA', 'Fiat\n\nAuto SA'),
    ('FCA de Ahorro para fines determinados SA', 'FCA de Ahorro para fines determinados SA'),
    ('Expreso Caraza SAC', 'Expreso Caraza SAC'),
])
def test_company_is_only_the_legal_name(monkeypatch, text, expected):
    monkeypatch.setattr(pipeline, 'ENABLE_PRESIDIO', False)
    monkeypatch.setattr(pipeline, 'ENABLE_SPACY', False)
    companies = [m for m in pipeline.run_detection(text) if m.cat == 'EMPRESA']
    assert [m.surface for m in companies] == [expected]
    assert text[companies[0].start:companies[0].end] == expected


@pytest.mark.parametrize('surface', ['Cámara dictaminó a fojas 215', 'Cámara a sus efectos', 'Juzgador'])
def test_generic_roles_and_court_prose_are_not_institutions(surface):
    assert not is_valid_detection('ORGANISMO', surface)


def test_two_companies_do_not_become_one(monkeypatch):
    monkeypatch.setattr(pipeline, 'ENABLE_PRESIDIO', False)
    monkeypatch.setattr(pipeline, 'ENABLE_SPACY', False)
    text = 'Armada SA y FCA Ahorro SA.'
    mentions = pipeline.run_detection(text, ['EMPRESA'])
    assert [m.surface for m in mentions] == ['Armada SA', 'FCA Ahorro SA']
    assert anonymize_text(text, build_detections_from_mentions(mentions, 'cat')) == '[EMPRESA_1] y [EMPRESA_2].'


def test_line_wraps_use_same_company_placeholder():
    text = 'Fiat Auto SA. Fiat\n\nAuto SA.'
    mentions = pipeline.run_detection(text, ['EMPRESA'])
    detections = build_detections_from_mentions(mentions, 'cat')
    assert len(detections) == 1
    assert len(detections[0].positions) == 2


@pytest.mark.parametrize('text, expected', [
    ('Firmado por: MATILDE BALLERINI, JUEZ DE CÁMARA', 'MATILDE BALLERINI'),
    ('Firmado por: AUGUSTO DANZI BIAUS, PROSECRETARIO DE CAMARA', 'AUGUSTO DANZI BIAUS'),
    ('El señor Donati adquirió un vehículo.', 'Donati'),
    ('El señor\n\nDonati suscribió el contrato.', 'Donati'),
    ('"Llanos, Andrea\n\nLaura c/ Fiat Auto SA s/ ordinario"', 'Llanos, Andrea\n\nLaura'),
    ('"Laborde de Ognian, Ethel B. c/ Universal Assistance SA"', 'Laborde de Ognian, Ethel B'),
    ('"Fernández Noble, Julia Iris c/ Fiat Auto SA s/ ordinario"', 'Fernández Noble, Julia Iris'),
])
def test_explicit_name_context_recognizes_missing_catalog_entries(monkeypatch, text, expected):
    monkeypatch.setattr(pipeline, 'ENABLE_PRESIDIO', False)
    monkeypatch.setattr(pipeline, 'ENABLE_SPACY', False)
    mentions = pipeline.run_detection(text, ['PERSONA'])
    assert any(m.surface == expected for m in mentions), [m.surface for m in mentions]
    assert expected not in anonymize_text(text, build_detections_from_mentions(mentions, 'cat'))


def test_unknown_first_name_is_not_trimmed_as_prose():
    assert _clean_persona_surface('Adriel Perez', 30) == ('Adriel Perez', 30)


def test_disabled_category_cannot_reserve_person_span(monkeypatch):
    monkeypatch.setattr(pipeline, 'ENABLE_PRESIDIO', False)
    monkeypatch.setattr(pipeline, 'ENABLE_SPACY', False)
    text = 'Juan Perez SA'
    assert [m.cat for m in pipeline.run_detection(text)] == ['EMPRESA']
    assert [m.surface for m in pipeline.run_detection(text, ['PERSONA'])] == ['Juan Perez']


def test_product_context_does_not_blacklist_real_surname():
    text = 'Juan Toro compró un vehículo Fiat Toro.'
    mentions = pipeline.run_detection(text, ['PERSONA'])
    assert any(m.surface == 'Juan Toro' for m in mentions)
    assert not any(m.surface == 'Fiat Toro' for m in mentions)


def test_confirmed_cluster_does_not_reuse_unrelated_person_placeholder():
    mentions = [Mention(id='m1', cat='PERSONA', surface='Juan Perez', start=0, end=10),
                Mention(id='m2', cat='PERSONA', surface='JUAN PEREZ', start=11, end=21)]
    detections = [Detection(id=0, cat='PERSONA', original='Otro Nombre', placeholder='[PERSONA_1]'),
                  Detection(id=1, cat='PERSONA', original='Juan Perez', placeholder='[PERSONA_2]', mention_ids=['m1']),
                  Detection(id=2, cat='PERSONA', original='JUAN PEREZ', placeholder='[PERSONA_3]', mention_ids=['m2'])]
    cluster = Cluster(cluster_id='sug_1', cat='PERSONA', mention_ids=['m1', 'm2'], surfaces=['Juan Perez', 'JUAN PEREZ'])
    state = SessionState(session_id='placeholder-test', mentions=mentions, detections=detections, clusters=[cluster])
    confirmed = confirm_cluster(state, cluster.cluster_id)
    assert confirmed.placeholder == '[PERSONA_4]'
    assert state.detections[0].placeholder == '[PERSONA_1]'
    assert state.detections[1].placeholder == state.detections[2].placeholder == '[PERSONA_4]'


PARAGRAPHS = [
    'Cámara Nacional de Apelaciones en lo Comercial', 'Sala B',
    'Gerardo Alfredo Donati solicita que se rechace la demanda.',
    'Por Gerardo Alfredo Donati contra Armada SA y FCA de Ahorro para fines determinados SA.',
    'El señor Donati adquirió de Armada SA un vehículo Fiat Toro.',
    'La Cámara dictaminó a fojas 215.',
    'Corresponde confirmar la responsabilidad de FCA Ahorro SA.',
    'Firmado por: MATILDE BALLERINI, JUEZ DE CÁMARA',
    'Firmado por: AUGUSTO DANZI BIAUS, PROSECRETARIO DE CAMARA',
    'MATILDE E. BALLERINI', 'AUGUSTO DANZI BIAUS', 'PROSECRETARIO DE CÁMARA',
    '"Llanos, Andrea\n\nLaura c/ Fiat Auto SA s/ ordinario"',
    '"Laborde de Ognian, Ethel B. c/ Universal Assistance SA"',
]


@pytest.mark.parametrize('categories', [None, ['PERSONA']])
def test_real_layers_preview_word_pdf_and_csv(categories):
    document = Document()
    for paragraph in PARAGRAPHS:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    with TestClient(app) as client:
        uploaded = client.post('/api/upload', files={'file': ('prueba_categorias.docx', buffer.getvalue(),
                                'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        sid = uploaded.json()['session_id']
        analysis = client.post('/api/analyze', json={'session_id': sid, 'label_mode': 'cat', 'enabled_categories': categories})
        assert analysis.status_code == 200, analysis.text
        detections = analysis.json()['detections']
        assert not any(d['original'] == 'Fiat Toro' for d in detections)
        if categories is None:
            companies = [d['original'] for d in detections if d['cat'] == 'EMPRESA']
            assert 'Armada SA' in companies and 'FCA Ahorro SA' in companies
            assert all('responsabilidad' not in name.lower() and 'Donati' not in name for name in companies)
            assert any(d['original'] == PARAGRAPHS[0] for d in detections if d['cat'] == 'ORGANISMO')
        payload = {'session_id': sid, 'use_confirmed_only': False}
        text = client.post('/api/export/preview', json=payload).json()['text']
        assert 'solicita que se rechace la demanda' in text
        assert 'La Cámara dictaminó a fojas 215.' in text
        assert 'confirmar la responsabilidad de' in text
        assert 'Fiat Toro' in text
        for original in ('Gerardo Alfredo Donati', 'MATILDE', 'BALLERINI', 'AUGUSTO', 'DANZI', 'BIAUS',
                         'Donati', 'Llanos', 'Laura', 'Laborde', 'Ognian', 'Ethel'):
            assert original not in text, (original, text)
        for extension in ('docx', 'pdf'):
            response = client.post('/api/export/' + extension, json=payload)
            assert response.status_code == 200, response.text
            if extension == 'docx':
                exported = '\n'.join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)
            else:
                with pdfplumber.open(io.BytesIO(response.content)) as pdf:
                    exported = '\n'.join(page.extract_text() or '' for page in pdf.pages)
            assert 'MATILDE' not in exported and 'Donati' not in exported
            assert 'Fiat Toro' in exported
        csv = client.post('/api/export/csv', json=payload)
        # CSV intentionally contains original equivalences; verify corrected category boundaries.
        assert csv.status_code == 200, csv.text
        assert 'confirmar la responsabilidad' not in csv.text
        assert ';Fiat Toro;' not in csv.text
