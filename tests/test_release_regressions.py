"""Regressions found while reviewing the complete release tree."""
import csv
import io
import re
import threading
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.api import routes_analyze, routes_upload
from app.detection import pipeline, presidio_layer, spacy_layer
from app.detection.text_chunks import iter_text_chunks
from app.export.csv_export import build_csv_bytes
from app.extraction.docx import extract_docx
from app.main import app
from app.models.schemas import AnalyzeResponse, Detection
from app.models.store import store
from app.services.analysis_cancel import AnalysisCancelledError, check_cancel, clear_cancel, request_cancel


def session(text='Juan Perez compareció. DNI 20.123.456.'):
    sid = store.create()
    state = store.get(sid)
    state.doc_text = text
    return sid


def test_reanalysis_restores_default_categories(monkeypatch):
    monkeypatch.setattr(pipeline, 'ENABLE_SPACY', False)
    monkeypatch.setattr(pipeline, 'ENABLE_PRESIDIO', False)
    sid = session()
    with TestClient(app) as client:
        first = client.post('/api/analyze', json={'session_id': sid, 'enabled_categories': ['PERSONA']})
        assert first.status_code == 200
        assert not any(d['cat'] == 'DNI' for d in first.json()['detections'])
        second = client.post('/api/analyze', json={'session_id': sid, 'enabled_categories': None})
        assert second.status_code == 200
        assert any(d['cat'] == 'DNI' for d in second.json()['detections'])
    store.delete(sid)


def test_empty_category_selection_leaves_existing_review_intact():
    sid = session()
    state = store.get(sid)
    state.detections = [Detection(id=0, cat='PERSONA', original='Juan Perez', placeholder='[PERSONA_8]')]
    with TestClient(app) as client:
        result = client.post('/api/analyze', json={'session_id': sid, 'enabled_categories': []})
    assert result.status_code == 400
    assert state.detections[0].placeholder == '[PERSONA_8]'
    store.delete(sid)


def test_cancel_endpoint_is_serviced_while_analysis_is_running(monkeypatch):
    started, release = threading.Event(), threading.Event()
    sid = session()

    def working(state):
        clear_cancel(sid)
        started.set()
        assert release.wait(10), 'El test no liberó el análisis.'
        check_cancel(sid)
        return AnalyzeResponse(session_id=sid, detections=[], clusters=[], stats={})

    monkeypatch.setattr(routes_analyze, 'run_full_analysis', working)
    with TestClient(app) as client, ThreadPoolExecutor(max_workers=2) as pool:
        analysis = pool.submit(client.post, '/api/analyze', json={'session_id': sid})
        try:
            assert started.wait(5)
            cancelled = pool.submit(client.post, '/api/analyze/cancel', json={'session_id': sid})
            assert cancelled.result(timeout=3).status_code == 200
        finally:
            release.set()
        assert analysis.result(timeout=5).status_code == 409
    store.delete(sid)


@pytest.mark.parametrize('length', [0, 1, 100_000, 100_001, 600_200])
def test_nlp_chunks_cover_all_original_characters(length):
    text = ('abcd ' * ((length + 4) // 5))[:length]
    covered = 0
    for offset, chunk in iter_text_chunks(text):
        assert text[offset:offset + len(chunk)] == chunk
        assert len(chunk) <= 100_000
        assert offset <= covered
        covered = offset + len(chunk)
    assert covered == length


def test_spacy_reads_person_after_old_500k_cutoff_with_original_offsets(monkeypatch):
    calls = []

    def nlp(chunk):
        calls.append(len(chunk))
        entities = [SimpleNamespace(label_='PER', text=m.group(), start_char=m.start(), end_char=m.end())
                    for m in re.finditer('Mariana Zorvik', chunk)]
        return SimpleNamespace(ents=entities)

    monkeypatch.setattr(spacy_layer, '_get_nlp', lambda: nlp)
    text = ('texto neutro. ' * 45_000) + 'Mariana Zorvik compareció.'
    matches = spacy_layer.detect_spacy(text)
    assert any(m.original == 'Mariana Zorvik' and m.start == text.index('Mariana Zorvik') for m in matches)
    assert len(calls) > 5 and max(calls) <= 100_000


def test_spacy_overlap_retains_a_name_split_at_the_chunk_edge(monkeypatch):
    def nlp(chunk):
        return SimpleNamespace(ents=[SimpleNamespace(label_='PER', text=m.group(), start_char=m.start(), end_char=m.end())
                                     for m in re.finditer('Mariana Zorvik', chunk)])
    monkeypatch.setattr(spacy_layer, '_get_nlp', lambda: nlp)
    text = ('x ' * 49_998) + 'Mariana Zorvik compareció.'
    matches = spacy_layer.detect_spacy(text)
    assert any(m.start == 99_996 and text[m.start:m.end] == 'Mariana Zorvik' for m in matches)


@pytest.mark.parametrize('layer', ['spacy', 'presidio'])
def test_cancellation_between_nlp_chunks_is_not_swallowed(monkeypatch, layer):
    sid = 'chunk-cancellation-' + layer
    calls = []

    def process(chunk):
        calls.append(chunk)
        request_cancel(sid)
        return []

    monkeypatch.setattr(pipeline, 'ENABLE_SPACY', layer == 'spacy')
    monkeypatch.setattr(pipeline, 'ENABLE_PRESIDIO', layer == 'presidio')
    if layer == 'spacy':
        monkeypatch.setattr(spacy_layer, '_get_nlp', lambda: lambda chunk: SimpleNamespace(ents=process(chunk)))
    else:
        monkeypatch.setattr(presidio_layer, '_get_analyzer', lambda: SimpleNamespace(analyze=lambda text, **kwargs: process(text)))
    with pytest.raises(AnalysisCancelledError):
        pipeline.run_detection('texto neutro. ' * 10_000, ['PERSONA'], session_id=sid)
    assert len(calls) == 1
    clear_cancel(sid)


def word_bytes(doc):
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_word_keeps_tables_in_body_order_and_reads_headers_and_footers():
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = 'IRINA ZORVIK'
    doc.add_paragraph('Inicio del documento.')
    row = doc.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = 'Mariana Quirna'
    row.cells[1].text = 'DNI 20.123.456'
    doc.add_paragraph('Cierre del documento.')
    doc.sections[0].footer.paragraphs[0].text = 'Matías Nolvik'
    text = extract_docx(word_bytes(doc))
    assert all(value in text for value in ('IRINA ZORVIK', 'Mariana Quirna', 'DNI 20.123.456', 'Matías Nolvik'))
    assert text.index('Inicio') < text.index('Mariana') < text.index('Cierre')
    assert 'Mariana Quirna | DNI' in text


def test_word_merged_cells_are_not_repeated_and_nested_tables_are_read():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = 'Irina Zorvik'
    nested = table.cell(1, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = 'Mariana Quirna'
    text = extract_docx(word_bytes(doc))
    assert text.count('Irina Zorvik') == 1
    assert 'Mariana Quirna' in text


def test_unsupported_upload_is_a_client_error():
    with TestClient(app) as client:
        response = client.post('/api/upload', files={'file': ('documento.txt', b'texto')})
    assert response.status_code == 400


def test_upload_limit_is_enforced_before_extraction(monkeypatch):
    monkeypatch.setattr(routes_upload, 'MAX_UPLOAD_BYTES', 8)
    with TestClient(app) as client:
        response = client.post('/api/upload', files={'file': ('grande.docx', b'x' * 20)})
    assert response.status_code == 413


@pytest.mark.parametrize('value', ['=2+2', '  +2+2', '\t-2+2', '\r@SUM(A1)', '\n =2+2'])
def test_csv_neutralizes_formulas_with_leading_whitespace(value):
    detections = [Detection(id=0, cat='OTRO', original=value, placeholder=value)]
    rows = list(csv.reader(io.StringIO(build_csv_bytes(detections).decode('utf-8-sig')), delimiter=';'))
    assert rows[1][1] == "'" + value and rows[1][2] == "'" + value
