import io

import pdfplumber
from docx import Document
from fastapi.testclient import TestClient

from app.anonymize.apply import anonymize_text
from app.main import app
from app.models.schemas import Detection, Position
from app.services.analyze import _prune_detections


def test_overlap_masks_union_once():
    text = 'Juan Perez Gomez declaró.'
    a = Detection(id=0, cat='PERSONA', original='Juan Perez', placeholder='[PERSONA_1]',
                  positions=[Position(start=0, end=10, raw='Juan Perez')])
    b = Detection(id=1, cat='PERSONA', original='Perez Gomez', placeholder='[PERSONA_2]',
                  positions=[Position(start=5, end=16, raw='Perez Gomez')])
    assert anonymize_text(text, [a, b]) == '[PERSONA_1] declaró.'


def test_export_pruning_preserves_session_ids_and_positions():
    full = Detection(id=12, cat='PERSONA', original='Juan Perez', placeholder='[PERSONA_1]',
                     positions=[Position(start=0, end=10, raw='Juan Perez')])
    part = Detection(id=42, cat='PERSONA', original='Perez', placeholder='[PERSONA_2]',
                     positions=[Position(start=5, end=10, raw='Perez'), Position(start=12, end=17, raw='Perez')])
    result = _prune_detections([full, part], 'Juan Perez. Perez declaró.')
    assert [d.id for d in result] == [12, 42]
    assert len(result[1].positions) == 1
    assert len(part.positions) == 2
    assert part.id == 42


def test_real_nlp_and_document_exports():
    paragraphs = [
        'Dr. Juan Perez solicita que se rechace la demanda.',
        'JUAN PÉREZ presentó un escrito.',
        'MARÍA JOSÉ GONZÁLEZ declaró como testigo.',
        'PÉREZ, JUAN CARLOS compareció.',
        'Sr. PÉREZ manifestó su voluntad.',
        'PODER JUDICIAL DE LA NACIÓN',
        'JUZGADO NACIONAL EN LO CIVIL',
    ]
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    with TestClient(app) as client:
        health = client.get('/health').json()
        for name in ('spacy', 'presidio'):
            assert health['nlp_layers'][name]['available'], health
        for name in ('nombres', 'apellidos', 'regex'):
            assert health['nlp_layers']['dictionaries'][name]['available'], health
        uploaded = client.post('/api/upload', files={'file': ('Juan Perez DNI 12345678.docx', buffer.getvalue(),
                                'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        assert uploaded.status_code == 200, uploaded.text
        sid = uploaded.json()['session_id']
        analysis = client.post('/api/analyze', json={'session_id': sid, 'label_mode': 'cat', 'enabled_categories': ['PERSONA']})
        assert analysis.status_code == 200, analysis.text
        surfaces = [d['original'] for d in analysis.json()['detections']]
        assert not any('solicita' in s.casefold() or 'rechace' in s.casefold() for s in surfaces)
        payload = {'session_id': sid, 'use_confirmed_only': False}
        preview = client.post('/api/export/preview', json=payload)
        assert preview.status_code == 200, preview.text
        text = preview.json()['text']
        for personal in ('Perez', 'PÉREZ', 'JUAN', 'MARÍA', 'GONZÁLEZ'):
            assert personal not in text, text
        assert 'solicita que se rechace la demanda.' in text
        assert 'PODER JUDICIAL DE LA NACIÓN' in text
        for extension in ('docx', 'pdf'):
            exported = client.post('/api/export/' + extension, json={**payload, 'text': text})
            assert exported.status_code == 200, exported.text
            assert '12345678' not in exported.headers['content-disposition']
            if extension == 'docx':
                extracted = '\n'.join(p.text for p in Document(io.BytesIO(exported.content)).paragraphs)
            else:
                with pdfplumber.open(io.BytesIO(exported.content)) as pdf:
                    extracted = '\n'.join(p.extract_text() or '' for p in pdf.pages)
            assert 'solicita' in extracted
            assert 'PÉREZ' not in extracted and 'Perez' not in extracted
        csv = client.post('/api/export/csv', json=payload)
        assert csv.status_code == 200
        assert 'PERSONA' in csv.text
