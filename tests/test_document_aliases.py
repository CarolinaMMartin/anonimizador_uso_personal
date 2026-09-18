"""Document-learned names: coverage, ambiguity, boundaries and exports."""
import io
import unicodedata

import pdfplumber
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.anonymize.apply import anonymize_text
from app.detection import pipeline
from app.detection.document_aliases import detect_document_aliases
from app.detection.regex_ar import RawItem
from app.main import app
from app.models.schemas import SessionState
from app.services.analyze import run_full_analysis
from app.services.clusters import confirm_cluster


def seed(text, surface):
    start = text.index(surface)
    return RawItem('PERSONA', surface, start, start + len(surface))


@pytest.mark.parametrize('full, short', [
    ('Eliana Zorvik', 'Zorvik'),
    ('Eliana Zorvik', 'ZORVIK'),
    ('Eliana Zorvik', 'ZorVIK'),
    ('Zorvik, Eliana', 'Zorvik'),
    ('Zorvik, Eliana', 'Eliana Zorvik'),
    ('Eliana Zorvik', 'Eliana'),
    ('Eliana Maria Zorvik', 'Eliana Zorvik'),
    ('Eliana Maria Zorvik', 'E. M. Zorvik'),
    ('Vorst Zelkin, Irina', 'Vorst Zelkin'),
    ('Irina de la Quirna', 'Quirna'),
    ('Ágata Sólvik', 'SOLVIK'),
    ('Agata Solvik', unicodedata.normalize('NFD', 'Sólvik')),
    ('Luan O’Kelan', "O'Kelan"),
    ('Fenna Ros-Vaak', 'Ros Vaak'),
    ('Li Qorvik', 'Li'),
    ('Мария Зорвик', 'Зорвик'),
    ('明 李', '李'),
    ('Eliana Zorvik', 'zorvik'),
    ('Irina de la Quirna', 'I. de la Quirna'),
    ('Irina Vorst y Zelkin', 'Vorst y Zelkin'),
    ('Irina Vorst y Zelkin', 'I. Vorst y Zelkin'),
])
def test_learned_components_have_exact_original_offsets(full, short):
    text = f'{full} compareció. La Dra. {short} indicó lo sucedido.'
    aliases = detect_document_aliases(text, [seed(text, full)], [])
    assert any(alias.original == short and alias.start == text.index(short, len(full)) for alias in aliases), [a.original for a in aliases]
    assert all(text[a.start:a.end] == a.original for a in aliases)
    assert not any('indicó' in a.original for a in aliases)


@pytest.mark.parametrize('tail', [
    'Elianazorvik declaró.',
    'Eliana, Zorvik y otras personas.',
    'Zorvik123 es un código.',
    'preZorvik es una palabra.',
    'prefix_Zorvik es una clave.',
    'Eliana dijo que Zorvik declaró.',
    'eliana se pronuncia así.',
    'zorvik figura en un glosario.',
])
def test_does_not_expand_into_prose_codes_or_lists(tail):
    text = f'Eliana Zorvik compareció. {tail}'
    aliases = detect_document_aliases(text, [seed(text, 'Eliana Zorvik')], [])
    in_tail = [a.original for a in aliases if a.start >= text.index(tail)]
    assert tail not in in_tail
    assert 'Eliana, Zorvik' not in in_tail
    assert 'Eliana dijo que Zorvik' not in in_tail
    if tail.startswith(('Elianazorvik', 'Zorvik123', 'preZorvik', 'prefix_', 'eliana', 'zorvik')):
        assert in_tail == []


def test_components_from_different_people_do_not_form_a_name():
    text = 'Eliana Zorvik compareció. Irina Vorst declaró. Eliana Vorst opinó.'
    aliases = detect_document_aliases(text, [seed(text, 'Eliana Zorvik'), seed(text, 'Irina Vorst')], [])
    assert 'Eliana Vorst' not in [alias.original for alias in aliases]


@pytest.mark.parametrize('surface', ['Eliana y Zorvik', 'Eliana de Zorvik', 'Eliana de la Zorvik'])
def test_unseen_particles_are_not_absorbed_from_prose(surface):
    text = f'Eliana Zorvik compareció. {surface} declararon.'
    aliases = detect_document_aliases(text, [seed(text, 'Eliana Zorvik')], [])
    assert surface not in [alias.original for alias in aliases]
    assert {'Eliana', 'Zorvik'} <= {alias.original for alias in aliases}


def test_particles_must_belong_to_the_same_seed():
    text = 'Eliana Zorvik declaró. Irina de Zorvik declaró. Eliana de Zorvik opinaron.'
    aliases = detect_document_aliases(text, [seed(text, 'Eliana Zorvik'), seed(text, 'Irina de Zorvik')], [])
    assert 'Eliana de Zorvik' not in [alias.original for alias in aliases]


def test_no_recursive_learning_from_aliases():
    text = 'E. Zorvik declaró. Zorvik indicó lo sucedido.'
    assert detect_document_aliases(text, [RawItem('PERSONA', 'Zorvik', 1, 7)], []) == []


@pytest.mark.parametrize('full', ['Mariana Vorst Zelkin', 'Matías de la Quirna Zelkin'])
def test_natural_compound_surname_groups_when_given_names_supply_boundary(monkeypatch, full):
    monkeypatch.setattr(pipeline, 'ENABLE_SPACY', False)
    monkeypatch.setattr(pipeline, 'ENABLE_PRESIDIO', False)
    surname = 'Vorst Zelkin' if 'Vorst' in full else 'Quirna Zelkin'
    state = SessionState(session_id='synthetic', doc_text=f'Dr. {full} compareció. {surname} indicó lo sucedido.')
    analysis = run_full_analysis(state)
    compound = next(d for d in analysis.detections if d.original == surname)
    assert compound.cluster_id is not None
    cluster = next(c for c in analysis.clusters if c.cluster_id == compound.cluster_id)
    assert full in cluster.surfaces


def test_protected_entities_are_not_learned_people():
    text = 'Eliana Zorvik compareció. Zorvik SA tiene sede en Calle Zorvik 123.'
    company, address = 'Zorvik SA', 'Calle Zorvik 123'
    protected = [RawItem(cat, value, text.index(value), text.index(value) + len(value))
                 for cat, value in [('EMPRESA', company), ('DOMICILIO', address)]]
    aliases = detect_document_aliases(text, [seed(text, 'Eliana Zorvik')], protected)
    assert all(alias.start < text.index(company) for alias in aliases)


@pytest.mark.parametrize('categories', [['PERSONA'], ['PERSONA', 'EMPRESA', 'DOMICILIO']])
def test_pipeline_and_unconfirmed_exports_keep_unknown_surname(monkeypatch, categories):
    monkeypatch.setattr(pipeline, 'ENABLE_SPACY', False)
    monkeypatch.setattr(pipeline, 'ENABLE_PRESIDIO', False)
    text = 'Matías Zorvik compareció. Más tarde, Zorvik asumió compromisos. ZORVIK indicó lo sucedido.'
    state = SessionState(session_id='synthetic', doc_text=text, enabled_categories=categories)
    analysis = run_full_analysis(state)
    short = next(d for d in analysis.detections if d.original == 'Zorvik')
    assert len(short.positions) == 2
    assert 'document_alias' in short.source_layers
    assert short.cluster_id is not None and not short.cluster_confirmed
    assert 'Zorvik' not in anonymize_text(text, state.detections)
    cluster = next(c for c in state.clusters if c.cluster_id == short.cluster_id)
    confirm_cluster(state, cluster.cluster_id)
    assert len({d.placeholder for d in state.detections if d.cat == 'PERSONA'}) == 1


def test_catalog_empty_still_reuses_explicit_unknown_full_name(monkeypatch):
    from app.detection import dictionaries
    monkeypatch.setattr(dictionaries, 'get_nombres', lambda: set())
    monkeypatch.setattr(dictionaries, 'get_apellidos', lambda: set())
    text = 'Firmado por: IRINA ZORVIK, JUEZA\n\nZorvik indicó lo sucedido.'
    aliases = detect_document_aliases(text, [seed(text, 'IRINA ZORVIK')], [])
    assert any(a.original == 'Zorvik' for a in aliases)


def test_ambiguous_surname_is_detected_but_does_not_merge_people(monkeypatch):
    monkeypatch.setattr(pipeline, 'ENABLE_SPACY', False)
    monkeypatch.setattr(pipeline, 'ENABLE_PRESIDIO', False)
    state = SessionState(session_id='synthetic', doc_text=(
        'Matías Zorvik compareció. Mariana Zorvik declaró. Zorvik indicó lo sucedido.'))
    analysis = run_full_analysis(state)
    assert any(d.original == 'Zorvik' for d in analysis.detections)
    assert not any('Matías Zorvik' in c.surfaces and 'Mariana Zorvik' in c.surfaces for c in analysis.clusters)
    assert 'Zorvik' not in anonymize_text(state.doc_text, state.detections)


def test_learned_vehicle_component_stays_visible(monkeypatch):
    monkeypatch.setattr(pipeline, 'ENABLE_SPACY', False)
    monkeypatch.setattr(pipeline, 'ENABLE_PRESIDIO', False)
    state = SessionState(session_id='synthetic', doc_text=(
        'Juan Toro declaró. Adquirió un vehículo Fiat Toro. Toro indicó lo sucedido.'), enabled_categories=['PERSONA'])
    analysis = run_full_analysis(state)
    preview = anonymize_text(state.doc_text, analysis.detections)
    assert 'Fiat Toro' in preview
    assert 'Toro indicó' not in preview


@pytest.mark.parametrize('categories', [['PERSONA'], None])
def test_real_docx_csv_word_pdf_cover_learned_alias(categories):
    doc = Document()
    for paragraph in ('MATÍAS ZORVIK declaró.', 'Zorvik asumió compromisos.', 'ZORVIK indicó lo sucedido.'):
        doc.add_paragraph(paragraph)
    stream = io.BytesIO()
    doc.save(stream)
    with TestClient(app) as client:
        uploaded = client.post('/api/upload', files={'file': ('ficticio.docx', stream.getvalue())})
        sid = uploaded.json()['session_id']
        response = client.post('/api/analyze', json={'session_id': sid, 'enabled_categories': categories})
        assert response.status_code == 200
        aliases = [d for d in response.json()['detections'] if d['cat'] == 'PERSONA' and d['original'] == 'Zorvik']
        assert len(aliases) == 1 and len(aliases[0]['positions']) == 2
        payload = {'session_id': sid}
        csv_export = client.post('/api/export/csv', json=payload)
        assert ';Zorvik;' in csv_export.content.decode('utf-8-sig')
        for extension in ('docx', 'pdf'):
            exported = client.post('/api/export/' + extension, json=payload)
            assert exported.status_code == 200
            if extension == 'docx':
                output = '\n'.join(p.text for p in Document(io.BytesIO(exported.content)).paragraphs)
            else:
                with pdfplumber.open(io.BytesIO(exported.content)) as pdf:
                    output = '\n'.join(page.extract_text() or '' for page in pdf.pages)
            assert 'Zorvik' not in output and 'ZORVIK' not in output
            assert 'asumió' in output and 'indicó' in output
