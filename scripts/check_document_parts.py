"""Verify Word parts and cancellation against an actual compiled local server."""
import argparse
import concurrent.futures
import io
import json
import time
import urllib.error
import urllib.request
import uuid
from pathlib import Path
from docx import Document
import portable_smoke as http

ROOT = Path(__file__).resolve().parents[1]

def upload(data, name='verificacion_ficticia.docx'):
    boundary = 'parts-' + uuid.uuid4().hex
    body = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{name}"\r\nContent-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n'.encode()
            + data + f'\r\n--{boundary}--\r\n'.encode())
    return json.loads(http.request('/api/upload', body, {'Content-Type': 'multipart/form-data; boundary=' + boundary}))

def document_bytes(document):
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def verify_manual_review():
    """El caso informado debe pasar también contra el ejecutable de Windows."""
    document = Document()
    document.add_paragraph('Documento ficticio para verificar revisión manual.')
    document.add_paragraph('📄 Campana 39. Repite: Campana 39. Fin.')
    sid = upload(document_bytes(document))['session_id']
    original = json.loads(http.request(f'/api/preview?session_id={sid}&mode=orig'))['text']
    start = original.index('Campana 39')
    payload = {'session_id': sid, 'cat': 'DOMICILIO', 'original': 'Campana 39',
               'start': start, 'end': start + len('Campana 39')}
    added = http.post('/api/manual-detection', payload)['detection']
    patch = urllib.request.Request(http.BASE + f'/api/detections/{added["id"]}',
        data=json.dumps({'session_id': sid, 'enabled': False}).encode(),
        headers={'Content-Type': 'application/json'}, method='PATCH')
    with http.OPENER.open(patch, timeout=30) as response:
        assert response.status == 200
    readded = http.post('/api/manual-detection', payload)['detection']
    preview = http.post('/api/export/preview', {'session_id': sid})['text']
    checks = {
        'manual_address_accepted': added['cat'] == 'DOMICILIO' and len(added['positions']) == 2,
        'manual_reselection_reactivates': readded['enabled'] and readded['id'] == added['id'],
        'preview_masks_both_occurrences': 'Campana 39' not in preview and preview.count(added['placeholder']) == 2,
        'unicode_context_preserved': '📄 ' in preview and 'Fin.' in preview,
    }
    for extension in ('docx', 'pdf'):
        raw = http.request('/api/export/' + extension,
            json.dumps({'session_id': sid}).encode(), {'Content-Type': 'application/json'})
        if extension == 'docx':
            text = '\n'.join(p.text for p in Document(io.BytesIO(raw)).paragraphs)
        else:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(raw)) as pdf:
                text = '\n'.join(p.extract_text() or '' for p in pdf.pages)
        checks[extension + '_manual_export'] = ('Campana 39' not in text
            and text.count(added['placeholder']) == 2 and 'Fin.' in text)
    return {'mode': 'manual_review', 'checks': checks}

def verify():
    document = Document()
    document.add_paragraph('El Dr. Juan Perez presentó un escrito.')
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'La Dra. Maria Gonzalez declaró en la tabla.'
    document.add_paragraph('Se conserva el texto posterior a la tabla.')
    section = document.sections[0]
    section.header.paragraphs[0].text = 'La Dra. Laura Suarez figura en el encabezado.'
    section.footer.paragraphs[0].text = 'El Dr. Diego Ramirez figura en el pie.'
    data = document_bytes(document)
    results = [verify_manual_review()]
    for categories in (['PERSONA'], None):
        sid = upload(data)['session_id']
        analyzed = http.post('/api/analyze', {'session_id': sid, 'enabled_categories': categories})
        preview = http.post('/api/export/preview', {'session_id': sid, 'use_confirmed_only': False})['text']
        names = ('Juan Perez', 'Maria Gonzalez', 'Laura Suarez', 'Diego Ramirez')
        checks = {
            'all_word_parts_detected': all(any(name in p['raw'] for d in analyzed['detections']
                if d['cat'] == 'PERSONA' for p in d['positions']) for name in names),
            'all_parts_hidden_in_preview': all(name not in preview for name in names),
            'table_and_following_text_preserved': 'declaró en la tabla' in preview and 'texto posterior a la tabla' in preview,
            'header_and_footer_context_preserved': 'figura en el encabezado' in preview and 'figura en el pie' in preview,
        }
        payload = json.dumps({'session_id': sid, 'use_confirmed_only': False}).encode()
        raw = http.request('/api/export/docx', payload, {'Content-Type': 'application/json'})
        word = Document(io.BytesIO(raw))
        text = '\n'.join(p.text for p in word.paragraphs)
        checks['word_export_hides_all_parts'] = all(name not in text for name in names) and 'declaró en la tabla' in text
        import pdfplumber
        raw = http.request('/api/export/pdf', payload, {'Content-Type': 'application/json'})
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
        checks['pdf_export_hides_all_parts'] = all(name not in text for name in names) and 'declaró en la tabla' in text
        results.append({'categories': categories or 'ALL', 'checks': checks})
    # This large fictitious document keeps a worker busy during the cancel request.
    large = Document()
    large.add_paragraph(('Texto de prueba para verificar cancelación. ' * 15000) + '\nDr. Juan Perez declaró.')
    sid = upload(document_bytes(large))['session_id']
    def analyze():
        try:
            http.post('/api/analyze', {'session_id': sid, 'enabled_categories': ['PERSONA']})
            return 200
        except urllib.error.HTTPError as exc:
            return exc.code
    with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
        future = pool.submit(analyze)
        time.sleep(0.5)
        started = time.monotonic()
        cancelled = http.post('/api/analyze/cancel', {'session_id': sid})
        elapsed = time.monotonic() - started
        checks = {'cancel_served_while_worker_busy': elapsed < 3 and cancelled['cancelled'],
                  'worker_reports_cancelled': future.result(timeout=90) == 409}
        results.append({'mode': 'cancel', 'checks': checks, 'response_seconds': round(elapsed, 3)})
    return results

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--url', default='http://127.0.0.1:8787')
    args = parser.parse_args()
    http.BASE = args.url.rstrip('/')
    results = verify()
    directory = ROOT / 'build/validation'
    directory.mkdir(parents=True, exist_ok=True)
    (directory / 'document_parts_portable.json').write_text(json.dumps(results, ensure_ascii=True, indent=2), encoding='utf-8')
    print(json.dumps(results, ensure_ascii=True, indent=2), flush=True)
    if not all(all(result['checks'].values()) for result in results):
        raise AssertionError('Compiled document extraction or cancellation failed')

if __name__ == '__main__':
    main()
